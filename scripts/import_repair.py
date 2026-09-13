#!/usr/bin/env python3
"""
Import dữ liệu Repair từ file Excel (.xlsx) NHIỀU SHEET vào database, chạy
trực tiếp trong terminal.

FORMAT FILE:
    Mỗi sheet = 1 loại board. Trong mỗi sheet:
      - Dòng 1, cột A: "{board_code} - {board_name}" (vd "12070000153 - HV
        Core V2.8 - Core 60/80/120/150Kw" -> code="12070000153",
        name="HV Core V2.8 - Core 60/80/120/150Kw" - chỉ tách ở dấu " - "
        ĐẦU TIÊN, phần còn lại (có thể chứa thêm dấu "-") thuộc về tên).
      - Dòng 2: bỏ trống (spacer).
      - Dòng 3: header (STT, Date receive, Board ID (gốc), Test Result,
        Failure Cause, Disposition, After repair, Push FW, Charging Station
        Test, Detailed Remarks, Ticket ID, SN, Date onsite, Original
        Phenomenon, Station Code).
      - Dòng 4 trở đi: data.
    Sheet nào không đúng format này (thiếu dòng 1 dạng "code - tên", hoặc
    thiếu cột bắt buộc ở dòng 3) sẽ tự động BỊ BỎ QUA, không báo lỗi - dùng để
    lọc ra các sheet không phải data (vd "Quick start guide", "Total_report").

QUY TRÌNH:
    1. Với mỗi sheet hợp lệ: tạo Board mới nếu code đó CHƯA có trong bảng
       `boards` (không đụng vào Board đã tồn tại sẵn, không xóa Board nào).
    2. Import toàn bộ dòng data của sheet đó vào bảng `repairs`, gắn
       board_code = code lấy từ dòng 1 của sheet.
    3. Mỗi dòng được tự sinh 1 "code" định danh thiết bị riêng (qua
       Repair.generate_code) - coi mỗi dòng là 1 thiết bị vật lý khác nhau.

    KHÔNG xử lý ảnh nhúng (không có cột ảnh nào được import trong format này).

    Mode 'reset' XÓA TOÀN BỘ dữ liệu Repair cũ + ảnh cũ trong storage TRƯỚC
    KHI import lại từ đầu - KHÔNG đụng tới bảng `boards` (board là danh mục
    lâu dài, không xóa theo mỗi lần import).

KHÔNG chạy trực tiếp file này - luôn chạy qua ./run_import.sh, vì đó là nơi
hỏi xác nhận trước khi xóa dữ liệu. File này sẽ từ chối chạy nếu thiếu cờ
--yes khi dùng mode 'reset' (chỉ được set bởi run_import.sh sau khi người
dùng xác nhận).

Usage (nội bộ, do run_import.sh gọi):
    python3 scripts/import_repair.py /path/to/file.xlsx --mode insert
    python3 scripts/import_repair.py /path/to/file.xlsx --mode reset --yes
"""
import argparse
import shutil
import sys
from pathlib import Path
from datetime import datetime

import openpyxl

from app.core.database import SessionLocal
from app.models.board import Board
# Contractor và User không dùng trực tiếp trong script này, nhưng BẮT BUỘC
# phải import vì Repair model có relationship("Contractor", ...) và kế thừa
# AuditMixin (relationship("User", ...) cho created_by/updated_by). SQLAlchemy
# chỉ resolve được tên class dạng string nếu class đó đã được import (đăng ký
# vào mapper registry) ở đâu đó trong cùng process. Thiếu 1 trong 2 dòng này
# sẽ lỗi ngay khi mapper được configure (vd lúc gọi db.query(Repair) lần đầu).
from app.models.contractor import Contractor
from app.models.user import User
from app.models.repairs import Repair
from app.models import TestResult


# Map tên cột hiển thị ở DÒNG 3 trong Excel -> tên field trong model Repair.
# Cột "STT" không map vì không dùng (chỉ là số thứ tự trong sheet, không lưu DB).
HEADER_TO_FIELD = {
    "Date receive": "date_receive",
    "Board ID (gốc)": "board_id",
    "Test Result": "test_tool_result",
    "Failure Cause": "failure_cause",
    "Disposition": "disposition",
    "After repair": "after_repair_result",
    "Push FW": "puss_f",
    "Charging Station Test": "charging_station_test_status",
    "Detailed Remarks": "detailed_remarks",
    "Ticket ID": "ticket_id",
    "SN": "sn",
    "Date onsite": "date_onsite",
    "Original Phenomenon": "original_phenomenon",
    "Station Code": "station_code",
}

# Các cột bắt buộc phải có ở dòng 3 - sheet nào thiếu 1 trong các cột này sẽ
# bị coi là "không đúng format", tự động bỏ qua (không phải lỗi).
REQUIRED_COLUMNS = [
    "Date receive", "Board ID (gốc)", "Test Result", "Failure Cause",
    "Disposition", "After repair", "Ticket ID", "SN",
    "Date onsite", "Original Phenomenon", "Station Code",
    "Charging Station Test", "Detailed Remarks",
]

HEADER_ROW = 3       # dòng chứa tên cột
DATA_START_ROW = 4   # dòng đầu tiên chứa data

REPAIR_PHOTO_STORAGE_DIR = "storage/repair_photos"

PROGRESS_EVERY = 500

# User được gán làm "người tạo" cho mọi record import qua script này - vì
# script chạy độc lập ngoài Telegram bot, không có context user thật (không
# đi qua current_user_id như khi tạo record qua bot), nên AuditMixin sẽ luôn
# để created_by_id/updated_by_id = NULL nếu không set thủ công ở đây.
IMPORT_USER_ID = 1

MODE_INSERT = "insert"
MODE_RESET = "reset"


class RepairImporter:
    def __init__(self, file_path: str, mode: str, confirmed: bool = False):
        self.file_path = file_path
        self.mode = mode
        self.confirmed = confirmed
        self.warnings = []
        self.failed = []
        self.created = 0
        self.boards_created = 0
        self.sheets_skipped = []

    # ------------------------------------------------------------------
    # Entry point
    # ------------------------------------------------------------------

    def run(self):
        if self.mode == MODE_RESET and not self.confirmed:
            print(
                "❌ Từ chối chạy: mode 'reset' sẽ XÓA TOÀN BỘ dữ liệu Repair cũ + "
                "toàn bộ ảnh trong storage trước khi import. Phải xác nhận trước "
                "(chạy qua ./run_import.sh, KHÔNG chạy file .py này trực tiếp)."
            )
            sys.exit(1)

        wb = openpyxl.load_workbook(self.file_path, data_only=True)

        valid_sheets = []
        for ws in wb.worksheets:
            board_code, board_name = self._parse_sheet_title(ws)
            if not board_code:
                self.sheets_skipped.append(f"'{ws.title}': dòng 1 không đúng dạng 'code - tên'")
                continue

            headers = [str(c.value).strip() if c.value else "" for c in ws[HEADER_ROW]]
            missing = [col for col in REQUIRED_COLUMNS if col not in headers]
            if missing:
                self.sheets_skipped.append(
                    f"'{ws.title}': thiếu cột bắt buộc ở dòng {HEADER_ROW}: {', '.join(missing)}"
                )
                continue

            valid_sheets.append((ws, board_code, board_name, headers))

        if self.sheets_skipped:
            print(f"ℹ️  Bỏ qua {len(self.sheets_skipped)} sheet không đúng format:")
            for s in self.sheets_skipped:
                print(f"  - {s}")
            print()

        if not valid_sheets:
            print("❌ Không tìm thấy sheet nào đúng format để import.")
            sys.exit(1)

        mode_label = "XÓA HẾT + IMPORT LẠI" if self.mode == MODE_RESET else "INSERT MỚI (giữ data cũ)"
        print(f"Bắt đầu import {len(valid_sheets)} sheet từ '{self.file_path}' - chế độ: {mode_label} ...\n")

        processed = 0
        with SessionLocal() as db:
            if self.mode == MODE_RESET:
                self._reset_all(db)

            for ws, board_code, board_name, headers in valid_sheets:
                try:
                    self._ensure_board(db, board_code, board_name)
                except Exception as e:
                    self.failed.append(f"Sheet '{ws.title}': không tạo được board '{board_code}' - {e}")
                    continue

                col_idx = {name: headers.index(name) for name in headers if name}
                total_rows_in_sheet = ws.max_row - (DATA_START_ROW - 1)
                print(f"--- Sheet '{ws.title}' (board {board_code}) - {max(total_rows_in_sheet, 0)} dòng data ---")

                for row_num, row in enumerate(
                    ws.iter_rows(min_row=DATA_START_ROW, values_only=True), start=DATA_START_ROW
                ):
                    if all(v is None for v in row):
                        continue

                    row_dict = {
                        field_name: self._normalize_placeholder(row[col_idx[excel_col]])
                        for excel_col, field_name in HEADER_TO_FIELD.items()
                        if excel_col in col_idx
                    }
                    location = f"Sheet '{ws.title}' dòng {row_num}"
                    try:
                        date_receive, w = self._parse_date(row_dict.get("date_receive"))
                        if w:
                            self.warnings.append(f"{location}: date_receive - {w}")

                        date_onsite, w = self._parse_date(row_dict.get("date_onsite"))
                        if w:
                            self.warnings.append(f"{location}: date_onsite - {w}")

                        test_tool_result, w = self._parse_enum(row_dict.get("test_tool_result"))
                        if w:
                            self.warnings.append(f"{location}: test_tool_result - {w}")

                        after_repair_result, w = self._parse_enum(row_dict.get("after_repair_result"))
                        if w:
                            self.warnings.append(f"{location}: after_repair_result - {w}")

                        # Chuẩn hóa ticket_id - giá trị normalized (None nếu là
                        # placeholder như "n/a") mới là thứ nên lưu vào DB, không
                        # phải literal text gốc, vì ticket_id có UNIQUE index.
                        ticket_id_str = self._normalize_ticket_id(row_dict.get("ticket_id"))

                        # Mỗi dòng coi như 1 thiết bị vật lý khác nhau -> tự
                        # sinh code riêng theo board_code của sheet này.
                        device_code = Repair.generate_code(db, board_code)

                        board_id_value = row_dict.get("board_id") or board_name or board_code

                        repair = Repair(
                            date_receive=date_receive,
                            board_id=str(board_id_value).strip(),
                            board_code=board_code,
                            code=device_code,
                            test_tool_result=test_tool_result,
                            failure_cause=row_dict.get("failure_cause"),
                            disposition=row_dict.get("disposition"),
                            after_repair_result=after_repair_result,
                            ticket_id=ticket_id_str,
                            sn=row_dict.get("sn"),
                            date_onsite=date_onsite,
                            original_phenomenon=row_dict.get("original_phenomenon"),
                            station_code=row_dict.get("station_code"),
                            charging_station_test_status=row_dict.get("charging_station_test_status"),
                            detailed_remarks=row_dict.get("detailed_remarks"),
                            puss_f=row_dict.get("puss_f"),
                            created_by_id=IMPORT_USER_ID,
                            updated_by_id=IMPORT_USER_ID,
                        )
                        db.add(repair)
                        # Commit NGAY từng dòng - không gộp batch. Lý do: lỗi
                        # UNIQUE constraint (vd trùng ticket_id giữa các sheet)
                        # chỉ lộ ra lúc commit, không phải lúc tạo object Repair()
                        # hay db.add(). Nếu gộp batch, 1 dòng lỗi sẽ làm rollback
                        # + mất luôn hàng trăm dòng khác đã xử lý đúng trong cùng
                        # batch đó. Commit từng dòng đảm bảo lỗi bị cô lập đúng
                        # 1 dòng duy nhất, các dòng khác không bị ảnh hưởng.
                        db.commit()
                        self.created += 1
                    except Exception as e:
                        db.rollback()
                        self.failed.append(f"{location}: {e}")
                        continue

                    processed += 1
                    if processed % PROGRESS_EVERY == 0:
                        print(f"  ... đã xử lý {processed} dòng (toàn bộ file)")

        self._print_report()

    # ------------------------------------------------------------------
    # Board: đọc dòng 1 sheet, tạo Board nếu chưa có
    # ------------------------------------------------------------------

    @staticmethod
    def _parse_sheet_title(ws):
        """Đọc ô A1, dạng 'CODE - TÊN BOARD' -> (code, name). Chỉ tách ở dấu
        ' - ' ĐẦU TIÊN - phần còn lại (có thể chứa thêm dấu '-') thuộc về tên.
        Trả về (None, None) nếu không đúng định dạng."""
        raw = ws.cell(row=1, column=1).value
        if not raw:
            return None, None
        text = str(raw).strip()
        if " - " not in text:
            return None, None
        code, _, name = text.partition(" - ")
        code = code.strip()
        name = name.strip() or None
        if not code:
            return None, None
        return code, name

    def _ensure_board(self, db, code, name):
        """Lấy Board theo code, tạo mới nếu chưa có. KHÔNG ghi đè tên nếu
        board đã tồn tại sẵn (tránh làm mất chỉnh sửa thủ công qua /board edit).

        Commit NGAY (không chỉ flush) sau khi tạo - vì giờ các dòng Repair
        được commit từng dòng một; nếu chỉ flush(), board mới tạo này sẽ nằm
        chung transaction với dòng Repair đầu tiên của sheet, và nếu dòng đó
        lỗi + rollback thì board cũng bị cuốn theo mất, dù board không có gì
        sai cả."""
        board = db.query(Board).filter(Board.code == code).first()
        if board:
            return board
        board = Board(code=code, ten=name, created_by_id=IMPORT_USER_ID, updated_by_id=IMPORT_USER_ID)
        db.add(board)
        db.commit()
        self.boards_created += 1
        return board

    # ------------------------------------------------------------------
    # Reset toàn bộ Repair (KHÔNG đụng tới Board) trước khi import
    # ------------------------------------------------------------------

    @staticmethod
    def _reset_all(db):
        print("⚠️  Đang xóa toàn bộ dữ liệu Repair cũ trong database (giữ nguyên danh mục Board) ...")
        try:
            deleted_count = db.query(Repair).delete()
            db.commit()
        except Exception as e:
            db.rollback()
            print(f"❌ Xóa dữ liệu cũ trong DB thất bại, dừng import. Chi tiết: {e}")
            sys.exit(1)
        print(f"  Đã xóa {deleted_count} bản ghi trong bảng repairs.")

        print(f"⚠️  Đang xóa toàn bộ ảnh cũ trong '{REPAIR_PHOTO_STORAGE_DIR}' ...")
        storage_path = Path(REPAIR_PHOTO_STORAGE_DIR)
        try:
            if storage_path.exists():
                shutil.rmtree(storage_path)
            storage_path.mkdir(parents=True, exist_ok=True)
        except Exception as e:
            print(f"❌ Xóa storage ảnh cũ thất bại, dừng import. Chi tiết: {e}")
            sys.exit(1)
        print("  Đã xóa xong.\n")

    # ------------------------------------------------------------------
    # Báo cáo kết quả
    # ------------------------------------------------------------------

    def _print_report(self):
        print()
        print(f"✅ Import xong: {self.created} bản ghi Repair mới, {self.boards_created} Board mới được tạo.")

        if self.failed:
            print(f"\n❌ {len(self.failed)} dòng lỗi (bị bỏ qua hoàn toàn):")
            for line in self.failed:
                print(f"  {line}")

        if self.warnings:
            print(f"\n⚠️ {len(self.warnings)} cảnh báo (dữ liệu vẫn được lưu, nhưng thiếu 1 phần):")
            for line in self.warnings:
                print(f"  {line}")

    # ------------------------------------------------------------------
    # Parse helpers
    # ------------------------------------------------------------------

    _PLACEHOLDER_VALUES = {"n/a", "na", "-", "none", "null"}

    @classmethod
    def _normalize_placeholder(cls, value):
        """Coi các placeholder kiểu 'n/a', '-', 'none', 'null', chuỗi rỗng...
        là None. Áp dụng chung cho MỌI cột."""
        if value is None:
            return None
        if isinstance(value, str):
            text = value.strip()
            if not text or text.lower() in cls._PLACEHOLDER_VALUES:
                return None
            return text
        return value

    @classmethod
    def _normalize_ticket_id(cls, value):
        return cls._normalize_placeholder(value)

    _DATE_FORMATS = (
        "%m/%d/%Y",
        "%m/%d/%y",
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y/%m/%d",
        # Cell được lưu dạng TEXT (không phải Date thật) nên có kèm giờ,
        # thường luôn là 00:00:00 vì Excel export ra kiểu này khi cell gốc
        # là ngày nhưng bị convert qua text - rất phổ biến trong file thật.
        "%Y-%m-%d %H:%M:%S",
        "%m/%d/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M:%S",
    )

    @classmethod
    def _parse_date(cls, value):
        if value is None or value == "":
            return None, None
        if isinstance(value, datetime):
            return value.date(), None
        if hasattr(value, "year") and hasattr(value, "month") and hasattr(value, "day"):
            return value, None

        text = str(value).strip()
        for fmt in cls._DATE_FORMATS:
            try:
                return datetime.strptime(text, fmt).date(), None
            except ValueError:
                continue
        return None, f"giá trị '{text}' không khớp format ngày nào được hỗ trợ -> để trống"

    @staticmethod
    def _parse_enum(value):
        if not value:
            return None, None
        text = str(value).strip()
        try:
            return TestResult(text.upper()), None
        except ValueError:
            return None, f"giá trị '{text}' không khớp enum TestResult -> để trống"


def main():
    parser = argparse.ArgumentParser(description="Import dữ liệu Repair (nhiều sheet, mỗi sheet 1 board) từ file Excel vào database.")
    parser.add_argument("file", help="Đường dẫn tới file .xlsx / .xlsm cần import")
    parser.add_argument(
        "--mode",
        choices=[MODE_INSERT, MODE_RESET],
        required=True,
        help=(
            f"'{MODE_INSERT}': giữ data cũ, chỉ insert thêm dòng mới. "
            f"'{MODE_RESET}': XÓA TOÀN BỘ Repair + ảnh cũ trước khi import lại từ đầu "
            f"(yêu cầu thêm --yes). Board KHÔNG bị xóa ở cả 2 mode."
        ),
    )
    parser.add_argument(
        "--yes",
        action="store_true",
        help=(
            "Bắt buộc khi --mode reset: xác nhận đồng ý xóa toàn bộ dữ liệu Repair "
            "cũ + ảnh cũ trước khi import. Cờ này chỉ nên được set bởi run_import.sh "
            "sau khi đã hỏi người dùng - không tự thêm cờ này để bỏ qua xác nhận."
        ),
    )
    args = parser.parse_args()

    file_path = Path(args.file)
    if not file_path.exists():
        print(f"❌ Không tìm thấy file: {file_path}")
        sys.exit(1)
    if file_path.suffix.lower() not in (".xlsx", ".xlsm"):
        print("❌ Chỉ hỗ trợ file .xlsx / .xlsm")
        sys.exit(1)

    try:
        wb_check = openpyxl.load_workbook(str(file_path), data_only=True)
        wb_check.close()
    except Exception as e:
        print(f"❌ Không đọc được file Excel: {e}")
        sys.exit(1)

    importer = RepairImporter(str(file_path), mode=args.mode, confirmed=args.yes)
    importer.run()


if __name__ == "__main__":
    main()