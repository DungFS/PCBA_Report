# app/bot/commands/contractor_report_command.py
import uuid
from datetime import datetime
from pathlib import Path

from telebot.types import InlineKeyboardMarkup, InlineKeyboardButton
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from app.bot.base_command import BaseCommand
from app.core.database import SessionLocal
from app.models import TestResult
from app.models.board import Board
from app.models.contractor import Contractor
from app.models.repairs import Repair

REPORT_TMP_DIR = "storage/tmp_reports"
CB_PREFIX = "contractorreport"
YEAR_PICKER_COUNT = 5  # số năm gần nhất cho chọn
MONTH_NAMES_VI = [
    "Tháng 1", "Tháng 2", "Tháng 3", "Tháng 4", "Tháng 5", "Tháng 6",
    "Tháng 7", "Tháng 8", "Tháng 9", "Tháng 10", "Tháng 11", "Tháng 12",
]


class ContractorReportCommand(BaseCommand):
    def register(self):
        @self.register_handler(commands=['contractor_report'], public=True)
        def handle(message):
            self._show_period_menu(message.chat.id)

        @self.bot.callback_query_handler(
            func=lambda call: bool(call.data) and call.data.startswith(f"{CB_PREFIX}:")
        )
        @self.admin_only
        def handle_callback(call):
            try:
                self.bot.answer_callback_query(call.id)
            except Exception:
                pass
            self._route_callback(call)

    def _route_callback(self, call):
        chat_id = call.message.chat.id
        parts = call.data.split(":")  # ["contractorreport", "<action>", ...]
        action = parts[1] if len(parts) > 1 else None

        if action == "all":
            self._send_contractor_report(chat_id, period_start=None, period_end=None, label="Tính đến hiện tại")
        elif action == "yearmenu":
            self._show_year_picker(chat_id)
        elif action == "year":
            year = int(parts[2])
            self._show_month_picker(chat_id, year)
        elif action == "month":
            year, month = int(parts[2]), int(parts[3])
            start, end = self._month_range(year, month)
            self._send_contractor_report(chat_id, period_start=start, period_end=end, label=f"Tháng {month}/{year}")

    # ------------------------------------------------------------------
    # Menu chọn khoảng thời gian
    # ------------------------------------------------------------------

    def _show_period_menu(self, chat_id):
        markup = InlineKeyboardMarkup(row_width=1)
        markup.add(
            InlineKeyboardButton("🗓 Theo tháng", callback_data=f"{CB_PREFIX}:yearmenu"),
            InlineKeyboardButton("♾ Tính đến hiện tại", callback_data=f"{CB_PREFIX}:all"),
        )
        self.bot.send_message(chat_id, "📊 Báo cáo Nhà thầu - chọn khoảng thời gian:", reply_markup=markup)

    def _show_year_picker(self, chat_id):
        current_year = datetime.now().year
        markup = InlineKeyboardMarkup(row_width=3)
        buttons = [
            InlineKeyboardButton(str(y), callback_data=f"{CB_PREFIX}:year:{y}")
            for y in range(current_year, current_year - YEAR_PICKER_COUNT, -1)
        ]
        markup.add(*buttons)
        self.bot.send_message(chat_id, "📅 Chọn năm:", reply_markup=markup)

    def _show_month_picker(self, chat_id, year):
        markup = InlineKeyboardMarkup(row_width=3)
        buttons = [
            InlineKeyboardButton(MONTH_NAMES_VI[m - 1], callback_data=f"{CB_PREFIX}:month:{year}:{m}")
            for m in range(1, 13)
        ]
        markup.add(*buttons)
        self.bot.send_message(chat_id, f"🗓 Chọn tháng trong năm {year}:", reply_markup=markup)

    @staticmethod
    def _month_range(year, month):
        """Trả về (start, end) - start = đầu tháng 00:00, end = đầu tháng kế
        tiếp 00:00 (khoảng nửa mở [start, end))."""
        start = datetime(year, month, 1)
        if month == 12:
            end = datetime(year + 1, 1, 1)
        else:
            end = datetime(year, month + 1, 1)
        return start, end

    # ------------------------------------------------------------------
    # Sinh báo cáo
    # ------------------------------------------------------------------

    def _send_contractor_report(self, chat_id, period_start, period_end, label):
        """
        Báo cáo 4 cột: ASP | Board Type | Quantity | First test result.

        Định nghĩa: với mỗi thiết bị (code), lấy bản ghi Repair CÓ ID NHỎ NHẤT
        (tức lần sửa chữa đầu tiên trong hệ thống, id tăng dần theo thứ tự tạo).
        Nhà thầu (ASP) của lần sửa đầu tiên đó được coi là "phụ trách" thiết bị này.
        Dữ liệu được gộp theo CẶP (ASP, Board Type) - 1 ASP có thể xuất hiện
        nhiều dòng nếu đã từng sửa nhiều loại bo khác nhau.
          - "Quantity": đếm số thiết bị (của loại bo đó) mà ASP phụ trách lần đầu.
          - "First test result": trong số đó, bao nhiêu thiết bị có
            test_tool_result = PASS ngay ở lần đầu (không cần sửa lại).

        LỌC THEO KHOẢNG THỜI GIAN: áp dụng lên created_at của chính lần sửa
        chữa ĐẦU TIÊN của thiết bị - tức "thiết bị này được đưa vào hệ thống
        lần đầu trong khoảng thời gian nào", KHÔNG PHẢI lọc theo các lần sửa
        lại sau đó. Vd thiết bị sửa lần đầu tháng 8, sửa lại tháng 9, thì báo
        cáo "Tháng 9" SẼ KHÔNG tính thiết bị đó.
          - period_start/period_end = None, None -> không lọc (tính đến hiện tại).
          - Có giá trị -> lọc theo nửa khoảng [period_start, period_end).
        Record thiếu created_at (dữ liệu cũ import trước khi có audit) bị loại
        khỏi báo cáo theo tháng cụ thể (không xác định được thuộc tháng nào),
        nhưng vẫn được tính đầy đủ trong "Tính đến hiện tại".

        Thiết bị mà lần sửa đầu tiên chưa gán nhà thầu (contractor_id NULL)
        được gộp vào dòng "(Chưa gán nhà thầu)". Board Type lấy từ boards.ten
        (tên hiển thị), fallback về board_code nếu board đó chưa đặt tên.
        """
        with SessionLocal() as db:
            all_repairs = (
                db.query(Repair)
                .filter(Repair.code.isnot(None))
                .order_by(Repair.id.asc())
                .all()
            )
            contractors = db.query(Contractor).order_by(Contractor.id).all()
            boards = db.query(Board).all()

        if not all_repairs:
            self.bot.send_message(chat_id, "📊 Chưa có dữ liệu repair nào để báo cáo.")
            return

        # Lấy record đầu tiên (id nhỏ nhất) cho mỗi code TRÊN TOÀN BỘ dữ liệu
        # (chưa lọc theo thời gian) - để xác định đúng "lần sửa đầu tiên thật
        # sự" của từng thiết bị, không bị ảnh hưởng bởi việc lọc sau đó.
        first_repair_per_device = {}
        for r in all_repairs:
            if r.code not in first_repair_per_device:
                first_repair_per_device[r.code] = r

        filtered = self._filter_by_period(first_repair_per_device.values(), period_start, period_end)

        if not filtered:
            self.bot.send_message(
                chat_id,
                f"📊 Không có thiết bị nào phát sinh lần sửa đầu tiên trong khoảng \"{label}\"."
            )
            return

        stats = {}  # (contractor_id, board_code) -> {"device_count", "pass_count"}
        for r in filtered:
            key = (r.contractor_id, r.board_code)
            entry = stats.setdefault(key, {"device_count": 0, "pass_count": 0})
            entry["device_count"] += 1
            if r.test_tool_result == TestResult.PASS:
                entry["pass_count"] += 1

        contractor_names = {c.id: c.name for c in contractors}
        board_names = {b.code: (b.ten or b.code) for b in boards}

        rows = []
        for (contractor_id, board_code), entry in stats.items():
            name = contractor_names.get(contractor_id, "(Chưa gán nhà thầu)")
            board_label = board_names.get(board_code, board_code or "(Chưa rõ loại bo)")
            rows.append((name, board_label, entry["device_count"], entry["pass_count"]))

        # Số lượng thiết bị nhiều nhất lên trên cho dễ nhìn
        rows.sort(key=lambda row: row[2], reverse=True)

        file_path = self._build_xlsx(rows)
        try:
            with open(file_path, "rb") as f:
                self.bot.send_document(
                    chat_id, f,
                    caption=(
                        f"📊 Báo cáo mức độ sử dụng theo Nhà thầu - {label}\n"
                        f"(tính theo lần sửa chữa ĐẦU TIÊN của mỗi thiết bị)"
                    )
                )
        finally:
            try:
                file_path.unlink()
            except Exception:
                pass  # dọn file tạm - không chặn nếu xóa lỗi

    @staticmethod
    def _filter_by_period(repairs, period_start, period_end):
        if period_start is None:
            return list(repairs)  # "Tính đến hiện tại" - không lọc
        return [
            r for r in repairs
            if getattr(r, "created_at", None) and period_start <= r.created_at < period_end
        ]

    @staticmethod
    def _build_xlsx(rows) -> Path:
        name_col = "ASP"
        board_col = "Board Type"
        device_col = "Quantity"
        pass_col = "First test result"

        wb = Workbook()
        ws = wb.active
        ws.title = "Contractor Report"

        header_font = Font(name="Arial", bold=True, size=11, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="4472C4")
        normal_font = Font(name="Arial", size=11)
        center = Alignment(horizontal="center", vertical="center")
        left = Alignment(horizontal="left", vertical="center")

        headers = [name_col, board_col, device_col, pass_col]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center

        for row_idx, (name, board_label, device_count, pass_count) in enumerate(rows, start=2):
            ws.cell(row=row_idx, column=1, value=name).font = normal_font
            ws.cell(row=row_idx, column=1).alignment = left
            ws.cell(row=row_idx, column=2, value=board_label).font = normal_font
            ws.cell(row=row_idx, column=2).alignment = left
            ws.cell(row=row_idx, column=3, value=device_count).font = normal_font
            ws.cell(row=row_idx, column=3).alignment = center
            ws.cell(row=row_idx, column=4, value=pass_count).font = normal_font
            ws.cell(row=row_idx, column=4).alignment = center

        ws.column_dimensions[get_column_letter(1)].width = max(
            20, max((len(r[0]) for r in rows), default=10) + 4
        )
        ws.column_dimensions[get_column_letter(2)].width = max(
            18, max((len(r[1]) for r in rows), default=10) + 4
        )
        ws.column_dimensions[get_column_letter(3)].width = 12
        ws.column_dimensions[get_column_letter(4)].width = 18
        ws.freeze_panes = "A2"

        dir_path = Path(REPORT_TMP_DIR)
        dir_path.mkdir(parents=True, exist_ok=True)
        file_path = dir_path / f"contractor_report_{uuid.uuid4().hex[:8]}.xlsx"
        wb.save(file_path)

        return file_path