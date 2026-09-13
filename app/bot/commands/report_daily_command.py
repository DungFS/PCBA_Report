import uuid
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

from telebot.types import Message
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from app.bot.base_command import BaseCommand
from app.core.database import SessionLocal
from app.models.board import Board
from app.models.contractor import Contractor
from app.models.repairs import Repair

REPORT_TMP_DIR = "storage/tmp_reports"

class ReportDailyCommand(BaseCommand):
    def register(self):
        @self.register_handler(commands=['report_daily'], admin=True)
        def handle(message: Message):
            msg = self.bot.send_message(
                message.chat.id,
                "📊 Nhập khoảng thời gian bạn muốn báo cáo tổng hợp theo định dạng:\n"
                "<b>DD/MM/YYYY - DD/MM/YYYY</b>\n\n"
                "<i>Ví dụ: 01/09/2026 - 13/09/2026</i>",
                parse_mode="HTML"
            )
            self.bot.register_next_step_handler(msg, self.process_date_input)

    def process_date_input(self, message: Message):
        text = message.text.strip()
        
        if text.startswith("/"):
            self.bot.send_message(message.chat.id, "❌ Đã hủy thao tác lấy báo cáo.")
            return

        try:
            start_str, end_str = [x.strip() for x in text.split("-")]
            start_date = datetime.strptime(start_str, "%d/%m/%Y").date()
            end_date = datetime.strptime(end_str, "%d/%m/%Y").date()
        except ValueError:
            self.bot.send_message(
                message.chat.id,
                "❌ Sai định dạng. Vui lòng gọi lại lệnh /report_daily và nhập đúng định dạng DD/MM/YYYY - DD/MM/YYYY."
            )
            return

        if start_date > end_date:
            self.bot.send_message(message.chat.id, "❌ Ngày bắt đầu không được lớn hơn ngày kết thúc.")
            return

        self._generate_and_send_report(message.chat.id, start_date, end_date)

    def _generate_and_send_report(self, chat_id, start_date, end_date):
        end_date_filter = end_date + timedelta(days=1)
        
        with SessionLocal() as db:
            repairs = (
                db.query(Repair)
                .filter(
                    Repair.date_receive >= start_date,
                    Repair.date_receive < end_date_filter
                )
                .all()
            )
            boards = db.query(Board).all()
            contractors = db.query(Contractor).all()

        if not repairs:
            self.bot.send_message(
                chat_id,
                f"📊 Không có dữ liệu repair nào từ ngày {start_date.strftime('%d/%m/%Y')} đến {end_date.strftime('%d/%m/%Y')}."
            )
            return

        board_names = {b.code: (b.ten or b.code) for b in boards}
        contractor_names = {c.id: c.name for c in contractors}

        # Khởi tạo dict đếm số lượng với mặc định là 0 cho các trạng thái
        summary = defaultdict(lambda: {
            "tt_pass": 0, 
            "tt_fail": 0, 
            "tt_discard": 0,
            "ar_pass": 0, 
            "ar_fail": 0
        })

        for r in repairs:
            r_date = r.date_receive.date() if isinstance(r.date_receive, datetime) else r.date_receive
            c_name = contractor_names.get(r.contractor_id, "-") if r.contractor_id else None

            # Khóa nhóm chỉ bao gồm Ngày, Board Code và Tên nhà thầu
            key = (r_date, r.board_code, c_name)
            
            # Lấy chuỗi giá trị và chuyển về in thường để dễ so sánh
            # Chú ý: Nếu trong DB bạn lưu là 'OK' hay 'NG', hãy đổi chữ 'pass' và 'fail' bên dưới cho khớp
            t_tool = str(r.test_tool_result.value).lower() if r.test_tool_result else ""
            a_repair = str(r.after_repair_result.value).lower() if r.after_repair_result else ""

            # Cộng dồn số lượng tương ứng
            if t_tool == "pass":
                summary[key]["tt_pass"] += 1
            elif t_tool == "fail":
                summary[key]["tt_fail"] += 1
            elif t_tool == "discard":
                summary[key]["tt_discard"] += 1
                
            if a_repair == "pass":
                summary[key]["ar_pass"] += 1
            elif a_repair == "fail":
                summary[key]["ar_fail"] += 1

        raw_rows = []
        for key, counts in summary.items():
            r_date, b_code, c_name = key
            b_name = board_names.get(b_code, b_code)
            raw_rows.append({
                "date_val": r_date,
                "date_str": r_date.strftime("%d/%m/%Y") if r_date else None,
                "b_code": b_code,
                "b_name": b_name,
                "c_name": c_name,
                "tt_pass": counts["tt_pass"],
                "tt_fail": counts["tt_fail"],
                "tt_discard": counts["tt_discard"],
                "ar_pass": counts["ar_pass"],
                "ar_fail": counts["ar_fail"],
            })
        
        # Sắp xếp danh sách ưu tiên theo Ngày nhận (mới nhất) -> Board code
        raw_rows.sort(key=lambda x: (x["date_val"] or datetime.min.date(), x["b_code"] or ""), reverse=True)
        
        final_rows = [
            [
                r["date_str"], r["b_code"], r["b_name"], r["c_name"], 
                r["tt_pass"], r["tt_fail"], r["tt_discard"], r["ar_pass"], r["ar_fail"]
            ]
            for r in raw_rows
        ]

        columns = [
            "Date Receive", "Board Code", "Board Name", "Contractor", 
            "Test Tool Pass", "Test Tool Fail", "Discard", "After Repair Pass", "After Repair Fail"
        ]

        file_path = self._build_xlsx(columns, final_rows)
        
        try:
            with open(file_path, "rb") as f:
                self.bot.send_document(
                    chat_id, f,
                    caption=(
                        f"📊 <b>Báo cáo tổng hợp nhóm</b>\n"
                        f"Từ ngày: {start_date.strftime('%d/%m/%Y')}\n"
                        f"Đến ngày: {end_date.strftime('%d/%m/%Y')}"
                    ),
                    parse_mode="HTML"
                )
        finally:
            try:
                file_path.unlink()
            except Exception:
                pass

    @staticmethod
    def _build_xlsx(columns, rows) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.title = "Daily Summary Report"

        header_font = Font(name="Arial", bold=True, size=11, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="4472C4")
        normal_font = Font(name="Arial", size=10)
        center = Alignment(horizontal="center", vertical="center")
        left = Alignment(horizontal="left", vertical="center", wrap_text=True)

        for col_idx, h in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            
            if h in ["Board Name", "Contractor"]:
                ws.column_dimensions[get_column_letter(col_idx)].width = 30
            elif "Pass" in h or "Fail" in h or h == "Discard":
                ws.column_dimensions[get_column_letter(col_idx)].width = 18
            else:
                ws.column_dimensions[get_column_letter(col_idx)].width = 20

        for row_idx, row_values in enumerate(rows, start=2):
            for col_idx, value in enumerate(row_values, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = normal_font
                # Căn giữa cho Ngày, Code và 5 cột đếm số lượng (cột 1, 2, 5, 6, 7, 8, 9)
                if col_idx in (1, 2, 5, 6, 7, 8, 9):
                    cell.alignment = center
                else:
                    cell.alignment = left

        ws.freeze_panes = "A2"

        dir_path = Path(REPORT_TMP_DIR)
        dir_path.mkdir(parents=True, exist_ok=True)
        file_path = dir_path / f"daily_summary_report_{uuid.uuid4().hex[:8]}.xlsx"
        wb.save(file_path)

        return file_path